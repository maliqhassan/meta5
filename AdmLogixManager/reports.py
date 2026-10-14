"""Excel (.xlsx) and PDF report generation."""
import os
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

from db import EXPORT_DIR, fmt_date, money, parse_date

# ------------------------------------------------------------ amount in words
_ONES = ["", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine", "Ten",
         "Eleven", "Twelve", "Thirteen", "Fourteen", "Fifteen", "Sixteen", "Seventeen",
         "Eighteen", "Nineteen"]
_TENS = ["", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety"]


def _words(n):
    if n < 20:
        return _ONES[n]
    if n < 100:
        return _TENS[n // 10] + ("-" + _ONES[n % 10] if n % 10 else "")
    if n < 1000:
        return _ONES[n // 100] + " Hundred" + (" " + _words(n % 100) if n % 100 else "")
    for div, name in ((10 ** 9, "Billion"), (10 ** 6, "Million"), (1000, "Thousand")):
        if n >= div:
            return _words(n // div) + " " + name + (" " + _words(n % div) if n % div else "")
    return ""


def amount_in_words(amount, currency="Riyals", sub="Halalas"):
    amount = round(abs(float(amount or 0)), 2)
    whole = int(amount)
    frac = int(round((amount - whole) * 100))
    text = (_words(whole) or "Zero") + " " + currency
    if frac:
        text += f" and {_words(frac)} {sub if frac > 1 else sub.rstrip('s')}"
    return text + " Only"


# ------------------------------------------------------------ helpers
def out_path(name, ext):
    os.makedirs(EXPORT_DIR, exist_ok=True)
    safe = "".join(c if c.isalnum() or c in " -_&." else "_" for c in name).strip()
    return os.path.join(EXPORT_DIR, f"{safe}.{ext}")


THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
TOTAL_FILL = PatternFill("solid", fgColor="E8EEF6")


def write_sheet(ws, titles, summary, headers, rows, widths, num_cols=(), total=None,
                footer_lines=()):
    """Generic formatted sheet: titles, summary box (label, value), table, total row."""
    ncol = len(headers)
    r = 1
    for i, t in enumerate(titles):
        ws.cell(r, 1, t).font = Font(bold=True, size=14 if i == 0 else 11)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        ws.cell(r, 1).alignment = Alignment(horizontal="center")
        r += 1
    r += 1
    for label, value in summary:
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(1, ncol - 2))
        c = ws.cell(r, ncol - 1, value)
        c.font = Font(bold=True)
        if isinstance(value, (int, float)):
            c.number_format = "#,##0.00"
        r += 1
    if summary:
        r += 1
    head_row = r
    for ci, h in enumerate(headers, 1):
        c = ws.cell(r, ci, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    r += 1
    for row in rows:
        for ci, v in enumerate(row, 1):
            c = ws.cell(r, ci, v)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if (ci - 1) in num_cols and isinstance(v, (int, float)):
                c.number_format = "#,##0.00"
            if isinstance(v, (date, datetime)):
                c.number_format = "DD-MMM-YY"
                c.alignment = Alignment(horizontal="center", vertical="top")
        r += 1
    if total:
        for ci, v in enumerate(total, 1):
            c = ws.cell(r, ci, v)
            c.font = Font(bold=True)
            c.fill = TOTAL_FILL
            c.border = BORDER
            if isinstance(v, (int, float)):
                c.number_format = "#,##0.00"
        r += 1
    for line in footer_lines:
        r += 1
        ws.cell(r, 1, line).font = Font(bold=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(head_row + 1, 1)
    ws.page_setup.orientation = "landscape" if ncol > 7 else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"{head_row}:{head_row}"


_STYLES = {
    "title": ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=13, alignment=1, spaceAfter=4),
    "sub": ParagraphStyle("s", fontName="Helvetica-Bold", fontSize=10.5, alignment=1, spaceAfter=3),
    "cell": ParagraphStyle("c", fontName="Helvetica", fontSize=8.5, leading=10.5),
    "cellr": ParagraphStyle("cr", fontName="Helvetica", fontSize=8.5, leading=10.5, alignment=2),
    "cellc": ParagraphStyle("cc", fontName="Helvetica", fontSize=8.5, leading=10.5, alignment=1),
    "head": ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=8.5, leading=10.5,
                           alignment=1, textColor=colors.white),
    "bold": ParagraphStyle("b", fontName="Helvetica-Bold", fontSize=9.5, leading=12),
    "note": ParagraphStyle("n", fontName="Helvetica", fontSize=9, leading=12),
}


def _esc(v):
    return str(v).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")


def write_pdf(path, titles, summary, headers, rows, widths, num_cols=(), center_cols=(),
              total=None, footer_lines=(), landscape_mode=False, signature=True, file_label=""):
    pagesize = landscape(A4) if landscape_mode else A4
    doc = SimpleDocTemplate(path, pagesize=pagesize, leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=18 * mm)
    avail = pagesize[0] - 28 * mm
    scale = avail / sum(widths)
    cw = [w * scale for w in widths]
    story = []
    for i, t in enumerate(titles):
        story.append(Paragraph(_esc(t), _STYLES["title" if i == 0 else "sub"]))
    story.append(Spacer(1, 6))

    if summary:
        data = [[Paragraph(f"<b>{_esc(l)}</b>", _STYLES["cell"]),
                 Paragraph(f"<b>{money(v) if isinstance(v, (int, float)) else _esc(v)}</b>", _STYLES["cellr"])]
                for l, v in summary]
        t = Table(data, colWidths=[avail - 45 * mm, 45 * mm])
        t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, colors.black),
                               ("TOPPADDING", (0, 0), (-1, -1), 2),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
        story += [t, Spacer(1, 8)]

    def cell(v, ci):
        if ci in num_cols and isinstance(v, (int, float)):
            return Paragraph(money(v), _STYLES["cellr"])
        if isinstance(v, (date, datetime)):
            return Paragraph(v.strftime("%d-%b-%y"), _STYLES["cellc"])
        return Paragraph(_esc("" if v is None else v), _STYLES["cellc" if ci in center_cols else "cell"])

    data = [[Paragraph(_esc(h), _STYLES["head"]) for h in headers]]
    data += [[cell(v, ci) for ci, v in enumerate(row)] for row in rows]
    style = [("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#555555")),
             ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
             ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
             ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
    if total:
        data.append([Paragraph(f"<b>{money(v) if isinstance(v, (int, float)) else _esc(v or '')}</b>",
                               _STYLES["cellr" if isinstance(v, (int, float)) else "cell"])
                     for v in total])
        style.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E8EEF6")))
    t = Table(data, colWidths=cw, repeatRows=1)
    t.setStyle(TableStyle(style))
    story.append(t)
    for line in footer_lines:
        story += [Spacer(1, 6), Paragraph(_esc(line), _STYLES["bold"])]
    if signature:
        story.append(Spacer(1, 40))
        sig = Table([["Prepared By", "", "Approved By"],
                     ["____________________", "", "____________________"]],
                    colWidths=[avail / 3] * 3)
        sig.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"),
                                 ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]))
        story.append(sig)

    printed = datetime.now().strftime("%d-%b-%Y %H:%M")

    def on_page(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.drawString(14 * mm, 9 * mm, f"Printed: {printed}   {file_label}")
        canvas.drawRightString(pagesize[0] - 14 * mm, 9 * mm, f"Page: {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return path


# ------------------------------------------------------------ daily bank book
def _day_data(db, day_iso, method):
    s = db.day_summary(day_iso, method)
    rows = db.txns(day_iso, day_iso, method=method)
    return s, rows


def _day_titles(db, d, method):
    return [db.get("company_name"),
            f"{db.get('report_title')} for the Month of {d:%B %Y}",
            db.get("report_subtitle") if method == "Bank" else "Cash in Hand"]


def _day_summary_lines(s, rows):
    lines = [("Opening Balance (B/F)", s["opening"])]
    for r in rows:
        if r["type"] == "Income":
            lines.append((f"Received: {r['description']}", r["amount"]))
    lines += [("Total Expense (Today)", s["expense"]), ("Closing Balance", s["closing"])]
    return lines


DAY_HEAD = ["Sr. No.", "Date", "Description", "Category", "Amount (SAR)", "Remark"]
DAY_W = [7, 13, 44, 26, 14, 14]


def _day_rows(rows):
    out = []
    for i, r in enumerate([r for r in rows if r["type"] == "Expense"], 1):
        out.append([i, parse_date(r["date"]), r["description"], r["category"] or "",
                    r["amount"], r["remark"] or ""])
    return out


def export_day(db, day_iso, fmt="pdf", method="Bank"):
    d = parse_date(day_iso)
    s, rows = _day_data(db, day_iso, method)
    titles = _day_titles(db, d, method)
    summary = _day_summary_lines(s, rows)
    data = _day_rows(rows)
    total = ["", "", "Total Amount of Expense", "", s["expense"], ""]
    name = f"{d:%d-%b-%Y} {'Bank' if method == 'Bank' else 'Cash'} Expense"
    if fmt == "pdf":
        return write_pdf(out_path(name, "pdf"), titles, summary, DAY_HEAD, data, DAY_W,
                         num_cols={4}, center_cols={0}, total=total, file_label=name)
    wb = Workbook()
    ws = wb.active
    ws.title = f"{d:%d-%b-%Y}"
    write_sheet(ws, titles, summary, DAY_HEAD, data, DAY_W, num_cols={4}, total=total)
    path = out_path(name, "xlsx")
    wb.save(path)
    return path


def export_month_workbook(db, year, month, method="Bank"):
    """One sheet per day (newest first) + a summary sheet, like the old bank expense file."""
    first = date(year, month, 1)
    last = date(year + (month == 12), month % 12 + 1, 1)
    days = [r[0] for r in db.q("SELECT DISTINCT date FROM transactions WHERE date>=? AND date<? "
                               "AND method=? ORDER BY date DESC", (first.isoformat(), last.isoformat(), method))]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    summ = []
    for dy in sorted(days):
        s = db.day_summary(dy, method)
        summ.append([parse_date(dy), s["opening"], s["income"], s["expense"], s["closing"]])
    ms = {"inc": sum(r[2] for r in summ), "exp": sum(r[3] for r in summ)}
    write_sheet(ws, [db.get("company_name"), f"Monthly Summary - {first:%B %Y} ({method})"], [],
                ["Date", "Opening", "Income", "Expense", "Closing"], summ, [14, 16, 16, 16, 16],
                num_cols={1, 2, 3, 4}, total=["Total", "", ms["inc"], ms["exp"], ""])
    for dy in days:
        d = parse_date(dy)
        s, rows = _day_data(db, dy, method)
        sh = wb.create_sheet(f"{d:%d-%b-%Y}")
        write_sheet(sh, _day_titles(db, d, method), _day_summary_lines(s, rows), DAY_HEAD,
                    _day_rows(rows), DAY_W, num_cols={4},
                    total=["", "", "Total Amount of Expense", "", s["expense"], ""])
    path = out_path(f"{first:%b-%Y} {method} Expense Workbook", "xlsx")
    wb.save(path)
    return path


# ------------------------------------------------------------ transactions report
def export_txn_report(db, rows, title, subtitle, fmt="pdf"):
    head = ["Sr.", "Date", "Type", "Category", "Description", "Employee / Client", "Method",
            "Income", "Expense", "Remark"]
    data, ti, te = [], 0.0, 0.0
    for i, r in enumerate(rows, 1):
        inc = r["amount"] if r["type"] == "Income" else None
        exp = r["amount"] if r["type"] == "Expense" else None
        ti += inc or 0
        te += exp or 0
        data.append([i, parse_date(r["date"]), r["type"], r["category"] or "", r["description"],
                     r["employee_name"] or r["company_name"] or "", r["method"] or "",
                     inc, exp, r["remark"] or ""])
    total = ["", "", "", "", "TOTAL", "", "", ti, te, ""]
    cats = {}
    for r in rows:
        if r["type"] == "Expense":
            cats[r["category"] or "Uncategorised"] = cats.get(r["category"] or "Uncategorised", 0) + r["amount"]
    footer = [f"Total Income: {money(ti)}    Total Expense: {money(te)}    Net: {money(ti - te)}"]
    footer += [f"   {k}: {money(v)}" for k, v in sorted(cats.items(), key=lambda x: -x[1])]
    widths = [6, 11, 9, 18, 46, 22, 8, 14, 14, 16]
    titles = [db.get("company_name"), title, subtitle]
    if fmt == "pdf":
        return write_pdf(out_path(title, "pdf"), titles, [], head, data, widths, num_cols={7, 8},
                         center_cols={0, 2, 6}, total=total, footer_lines=["Expense by category:"] + footer,
                         landscape_mode=True, file_label=title)
    wb = Workbook()
    write_sheet(wb.active, titles, [], head, data, widths, num_cols={7, 8}, total=total,
                footer_lines=footer)
    wb.active.title = "Report"
    path = out_path(title, "xlsx")
    wb.save(path)
    return path


# ------------------------------------------------------------ employees
def export_employees(db, rows, fmt="pdf"):
    extras = db.custom_fields("employee")
    head = ["S.No", "ID No", "Iqama No", "Name", "Basic Salary", "Status", "Iqama Expiry",
            "Phone"] + extras
    data = []
    for i, r in enumerate(rows, 1):
        ex = db.extra(r)
        data.append([i, r["emp_code"], r["iqama_no"], r["name"], r["basic_salary"] or 0, r["status"],
                     fmt_date(r["iqama_expiry"]), r["phone"] or ""] + [ex.get(f, "") for f in extras])
    total = ["", "", "", "Total", sum(r[4] for r in data), "", "", ""] + [""] * len(extras)
    widths = [6, 9, 12, 38, 13, 11, 13, 14] + [14] * len(extras)
    title = f"Nominal Roll of Employees - {date.today():%B %Y}"
    titles = [db.get("company_name"), title]
    if fmt == "pdf":
        return write_pdf(out_path(title, "pdf"), titles, [], head, data, widths, num_cols={4},
                         center_cols={0, 1, 2, 5}, total=total, landscape_mode=len(head) > 8,
                         file_label=title)
    wb = Workbook()
    write_sheet(wb.active, titles, [], head, data, widths, num_cols={4}, total=total)
    wb.active.title = "Employees"
    path = out_path(title, "xlsx")
    wb.save(path)
    return path


# ------------------------------------------------------------ payroll
def export_payroll(db, month, fmt="pdf"):
    rows = db.payroll(month)
    m = datetime.strptime(month, "%Y-%m")
    head = ["S.No", "ID No", "Name", "Iqama No", "Basic", "Allow.", "Overtime", "Deduct.",
            "Advance", "Net Salary", "Paid", "Balance", "Paid Date", "Remark"]
    data = []
    for i, r in enumerate(rows, 1):
        net = db.net_salary(r)
        data.append([i, r["emp_code"], r["name"], r["iqama_no"], r["basic_salary"], r["allowance"],
                     r["overtime"], r["deduction"], r["advance"], net, r["paid_amount"],
                     net - (r["paid_amount"] or 0), fmt_date(r["paid_date"]), r["remark"] or ""])
    total = ["", "", "TOTAL", ""] + [sum((x[c] or 0) for x in data) for c in range(4, 12)] + ["", ""]
    widths = [5, 8, 34, 11, 10, 9, 9, 9, 9, 11, 11, 11, 11, 14]
    title = f"Salary Sheet - {m:%B %Y}"
    titles = [db.get("company_name"), title]
    nums = set(range(4, 12))
    if fmt == "pdf":
        return write_pdf(out_path(title, "pdf"), titles, [], head, data, widths, num_cols=nums,
                         center_cols={0, 1, 3, 12}, total=total, landscape_mode=True, file_label=title)
    wb = Workbook()
    write_sheet(wb.active, titles, [], head, data, widths, num_cols=nums, total=total)
    wb.active.title = f"{m:%b-%Y}"
    path = out_path(title, "xlsx")
    wb.save(path)
    return path


def export_payslip(db, payroll_id):
    r = db.one("SELECT p.*, e.emp_code, e.name, e.iqama_no FROM payroll p JOIN employees e "
               "ON e.id=p.employee_id WHERE p.id=?", (payroll_id,))
    m = datetime.strptime(r["month"], "%Y-%m")
    net = db.net_salary(r)
    rows = [["Basic Salary", r["basic_salary"]], ["Allowance", r["allowance"]],
            ["Overtime", r["overtime"]], ["Deduction", -(r["deduction"] or 0)],
            ["Advance Adjusted", -(r["advance"] or 0)]]
    summary = [("Employee", f"{r['emp_code']}  {r['name']}"), ("Iqama No", r["iqama_no"] or ""),
               ("Month", f"{m:%B %Y}"), ("Paid On", fmt_date(r["paid_date"]) or "-"),
               ("Payment Method", r["method"] or "")]
    name = f"Payslip {r['emp_code']} {m:%b-%Y}"
    return write_pdf(out_path(name, "pdf"), [db.get("company_name"), f"Salary Slip - {m:%B %Y}"],
                     summary, ["Particulars", "Amount (SAR)"], rows, [70, 30], num_cols={1},
                     total=["Net Salary", net],
                     footer_lines=[f"Paid: {money(r['paid_amount'])}",
                                   f"In words: {amount_in_words(r['paid_amount'] or net)}"],
                     file_label=name)


# ------------------------------------------------------------ client statement
def export_statement(db, company_id, fmt="pdf"):
    c = db.one("SELECT * FROM companies WHERE id=?", (company_id,))
    invs = db.invoices(company_id)
    head = ["Sr. #", "Invoice #", "Date", "Month / Period", "Description", "Amount", "Received",
            "Balance", "Remark"]
    data = []
    for i, r in enumerate(invs, 1):
        data.append([i, r["invoice_no"] or "", parse_date(r["issue_date"]), r["period"] or "",
                     r["description"] or "", r["amount"], r["received"], r["amount"] - r["received"],
                     r["remark"] or ""])
    t_amt = sum(r[5] for r in data)
    t_rec = sum(r[6] for r in data)
    due = t_amt - t_rec
    summary = [("Total Amount", t_amt), ("Received Amount", t_rec), ("Remaining Amount", due)]
    total = ["", "", "", "", "Total Amount (SAR)", t_amt, t_rec, due, ""]
    footer = ["To-date Due Amount in words:", amount_in_words(due)]
    titles = [c["name"], f"Statement of Account - Date: {date.today():%d-%b-%Y}",
              f"From: {db.get('company_name')}"]
    widths = [6, 14, 11, 22, 30, 13, 13, 13, 14]
    name = f"Statement {c['name']} {date.today():%d-%b-%Y}"
    if fmt == "pdf":
        return write_pdf(out_path(name, "pdf"), titles, summary, head, data, widths,
                         num_cols={5, 6, 7}, center_cols={0, 1}, total=total, footer_lines=footer,
                         landscape_mode=True, file_label=name)
    wb = Workbook()
    write_sheet(wb.active, titles, summary, head, data, widths, num_cols={5, 6, 7}, total=total,
                footer_lines=footer)
    wb.active.title = (c["code"] or "Statement")[:30]
    path = out_path(name, "xlsx")
    wb.save(path)
    return path


def export_companies(db, rows, fmt="pdf"):
    extras = db.custom_fields("company")
    head = ["S.No", "ID No", "Company Name", "Category", "Contact", "Phone", "VAT No",
            "Receivable"] + extras
    data = []
    for i, r in enumerate(rows, 1):
        due = db.scalar("SELECT SUM(amount) FROM invoices WHERE company_id=?", (r["id"],)) - db.scalar(
            "SELECT SUM(r.amount) FROM receipts r JOIN invoices i ON i.id=r.invoice_id WHERE i.company_id=?",
            (r["id"],))
        ex = db.extra(r)
        data.append([i, r["code"], r["name"], r["category"], r["contact_person"] or "", r["phone"] or "",
                     r["vat_no"] or "", due] + [ex.get(f, "") for f in extras])
    widths = [6, 8, 40, 16, 18, 14, 16, 14] + [14] * len(extras)
    title = "ID Wise Companies List"
    titles = [db.get("company_name"), title]
    total = ["", "", "Total", "", "", "", "", sum(r[7] for r in data)] + [""] * len(extras)
    if fmt == "pdf":
        return write_pdf(out_path(title, "pdf"), titles, [], head, data, widths, num_cols={7},
                         center_cols={0, 1}, total=total, landscape_mode=True, file_label=title)
    wb = Workbook()
    write_sheet(wb.active, titles, [], head, data, widths, num_cols={7}, total=total)
    wb.active.title = "Companies"
    path = out_path(title, "xlsx")
    wb.save(path)
    return path
